from __future__ import annotations

import hashlib
import json
from datetime import datetime, timezone
from io import BytesIO
from pathlib import Path
from typing import Any

import yaml
from fastapi import APIRouter, Depends, HTTPException, Query, status
from fastapi.responses import FileResponse

from backend.auth import authorize_cliente_access, get_current_user
from backend.repositories.file_repository import (
    append_hallazgo,
    read_hallazgos,
    read_memo,
    read_perfil,
    read_workpapers,
    write_memo,
)
from backend.routes.dashboard import get_dashboard
from backend.routes.workpapers import _generate_tasks, _merge_saved_tasks, _quality_gates
from backend.schemas import ApiResponse, PdfSummaryResponse, ReportMemoResponse, UserContext
from backend.schemas import (
    InternalControlLetterRequest,
    InternalControlLetterResponse,
    NIIFPymesDraftRequest,
    NIIFPymesDraftResponse,
)
from backend.services.rag_chat_service import generate_chat_response
from backend.services.report_generation_service import generate_internal_control_letter, generate_niif_pymes_draft

router = APIRouter(prefix="/reportes", tags=["reportes"])
ROOT = Path(__file__).resolve().parents[2]
EXPORTS_DIR = ROOT / "data" / "exports"
REPORTS_META_DIR = ROOT / "data" / "clientes"
DOC_REGISTRY_FILE = "documentos_registro.json"
DOC_LINKS_FILE = "document_section_links.json"
EVIDENCE_GATE_POLICIES_PATH = ROOT / "backend" / "templates" / "docx" / "evidence_gate_policies.yaml"
DOC_STATES = ["draft", "reviewed", "approved", "issued"]
ALLOWED_TRANSITIONS: dict[str, dict[str, list[str]]] = {
    "draft": {"reviewed": ["senior", "gerente", "socio"]},
    "reviewed": {"approved": ["gerente", "socio"]},
    "approved": {"issued": ["socio"]},
    "issued": {},
}
ROLE_ALIASES = {
    "staff": "staff",
    "senior": "senior",
    "gerente": "gerente",
    "manager": "gerente",
    "socio": "socio",
    "partner": "socio",
}
DEFAULT_EVIDENCE_POLICY: dict[str, Any] = {
    "enforce_evidence_gate_on_approved": False,
    "enforce_evidence_gate_on_issued": False,
    "minimum_coverage_percent": 0,
    "critical_sections": [],
}


def _safe_text(value: Any) -> str:
    return str(value or "").strip()


def _history_path(cliente_id: str) -> Path:
    return REPORTS_META_DIR / cliente_id / "reportes_historial.json"


def _registry_path(cliente_id: str) -> Path:
    return REPORTS_META_DIR / cliente_id / DOC_REGISTRY_FILE


def _links_path(cliente_id: str) -> Path:
    return REPORTS_META_DIR / cliente_id / DOC_LINKS_FILE


def _read_history(cliente_id: str) -> list[dict[str, Any]]:
    path = _history_path(cliente_id)
    if not path.exists():
        return []
    try:
        payload = json.loads(path.read_text(encoding="utf-8"))
    except Exception:
        return []
    if not isinstance(payload, list):
        return []
    return [x for x in payload if isinstance(x, dict)]


def _append_history(cliente_id: str, record: dict[str, Any]) -> None:
    history = _read_history(cliente_id)
    history.append(record)
    path = _history_path(cliente_id)
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(history[-50:], ensure_ascii=False, indent=2), encoding="utf-8")


def _read_registry(cliente_id: str) -> dict[str, Any]:
    path = _registry_path(cliente_id)
    if not path.exists():
        return {"documents": {}}
    try:
        payload = json.loads(path.read_text(encoding="utf-8"))
    except Exception:
        return {"documents": {}}
    if not isinstance(payload, dict):
        return {"documents": {}}
    docs = payload.get("documents")
    if not isinstance(docs, dict):
        payload["documents"] = {}
    return payload


def _write_registry(cliente_id: str, payload: dict[str, Any]) -> None:
    path = _registry_path(cliente_id)
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")


def _read_links(cliente_id: str) -> dict[str, Any]:
    path = _links_path(cliente_id)
    if not path.exists():
        return {"documents": {}}
    try:
        payload = json.loads(path.read_text(encoding="utf-8"))
    except Exception:
        return {"documents": {}}
    if not isinstance(payload, dict):
        return {"documents": {}}
    docs = payload.get("documents")
    if not isinstance(docs, dict):
        payload["documents"] = {}
    return payload


def _write_links(cliente_id: str, payload: dict[str, Any]) -> None:
    path = _links_path(cliente_id)
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")


def _document_type_for_kind(kind: str) -> str:
    mapping = {
        "internal_control_letter": "carta_control_interno",
        "niif_pymes_draft": "niif_pymes_borrador",
    }
    return mapping.get(kind, kind)


def _normalize_role(role: str) -> str:
    r = _safe_text(role).lower()
    return ROLE_ALIASES.get(r, r or "staff")


def _allowed_next_states(*, role: str, current_state: str) -> list[str]:
    role_n = _normalize_role(role)
    transitions = ALLOWED_TRANSITIONS.get(current_state, {})
    out: list[str] = []
    for target, roles in transitions.items():
        if role_n in roles:
            out.append(target)
    return out


def _read_evidence_gate_policies() -> dict[str, Any]:
    if not EVIDENCE_GATE_POLICIES_PATH.exists():
        return {}
    try:
        payload = yaml.safe_load(EVIDENCE_GATE_POLICIES_PATH.read_text(encoding="utf-8")) or {}
    except Exception:
        return {}
    return payload if isinstance(payload, dict) else {}


def _document_evidence_policy(document_type: str) -> dict[str, Any]:
    policies = _read_evidence_gate_policies()
    raw = policies.get(document_type) if isinstance(policies, dict) else None
    if not isinstance(raw, dict):
        return dict(DEFAULT_EVIDENCE_POLICY)
    critical = raw.get("critical_sections")
    critical_sections = [str(x).strip() for x in (critical if isinstance(critical, list) else []) if str(x).strip()]
    minimum = raw.get("minimum_coverage_percent")
    try:
        min_pct = float(minimum)
    except Exception:
        min_pct = 0.0
    min_pct = max(0.0, min(100.0, min_pct))
    return {
        "enforce_evidence_gate_on_approved": bool(raw.get("enforce_evidence_gate_on_approved", False)),
        "enforce_evidence_gate_on_issued": bool(raw.get("enforce_evidence_gate_on_issued", False)),
        "minimum_coverage_percent": min_pct,
        "critical_sections": critical_sections,
    }


def _section_matches_critical(section_id: str, critical_section: str) -> bool:
    sid = _safe_text(section_id)
    crit = _safe_text(critical_section)
    if not sid or not crit:
        return False
    if sid == crit or sid.startswith(f"{crit}:"):
        return True
    # Compatibilidad entre IDs legacy y nuevos.
    if crit == "findings" and (sid.startswith("hallazgo:") or sid.startswith("finding:")):
        return True
    if crit == "hallazgo" and (sid.startswith("hallazgo:") or sid.startswith("finding:")):
        return True
    return False


def _is_pending_marker(value: Any) -> bool:
    text = _safe_text(value).lower()
    return not text or "[[pendiente]]" in text


def _critical_hallazgo_ids(document_snapshot: dict[str, Any]) -> set[str]:
    out: set[str] = set()
    if not isinstance(document_snapshot, dict):
        return out
    findings = document_snapshot.get("findings")
    if not isinstance(findings, list):
        major = (
            document_snapshot.get("major_interest_findings")
            if isinstance(document_snapshot.get("major_interest_findings"), list)
            else []
        )
        control = (
            document_snapshot.get("internal_control_findings")
            if isinstance(document_snapshot.get("internal_control_findings"), list)
            else []
        )
        findings = [*major, *control]
    for idx, finding in enumerate(findings, start=1):
        if not isinstance(finding, dict):
            continue
        fid = _safe_text(finding.get("id")) or str(idx)
        prioridad = _safe_text(finding.get("prioridad")).lower()
        categoria = _safe_text(finding.get("categoria")).lower()
        if prioridad in {"alta", "critica", "crítica"} or categoria == "mayor_interes":
            out.add(fid)
    return out


def _coherence_issues_for_document(*, cliente_id: str, version: dict[str, Any], for_issue: bool) -> list[str]:
    issues: list[str] = []
    snapshot = version.get("document_snapshot") if isinstance(version.get("document_snapshot"), dict) else {}
    if not isinstance(snapshot, dict):
        return issues
    perfil = read_perfil(cliente_id) or {}
    encargo = perfil.get("encargo") if isinstance(perfil.get("encargo"), dict) else {}
    expected_period = _safe_text(encargo.get("periodo_fin"))
    cover = snapshot.get("cover") if isinstance(snapshot.get("cover"), dict) else {}
    header = snapshot.get("header") if isinstance(snapshot.get("header"), dict) else {}
    doc_period = _safe_text((cover or {}).get("period_end")) or _safe_text((header or {}).get("period_end"))
    if expected_period and doc_period and expected_period != doc_period:
        issues.append(f"Periodo inconsistente: documento '{doc_period}' != perfil '{expected_period}'.")
    if _safe_text(version.get("document_type")) == "carta_control_interno" and for_issue:
        ruc = _safe_text((cover or {}).get("ruc"))
        if _is_pending_marker(ruc):
            issues.append("RUC vacío o pendiente en carta a emitir.")
    return issues


def _section_hashes(document: dict[str, Any]) -> dict[str, str]:
    hashes: dict[str, str] = {}
    if not isinstance(document, dict):
        return hashes
    if isinstance(document.get("sections"), list):
        for section in document.get("sections", []):
            if not isinstance(section, dict):
                continue
            sid = _safe_text(section.get("id")) or _safe_text(section.get("title"))
            if not sid:
                continue
            raw = json.dumps(section, ensure_ascii=False, sort_keys=True)
            hashes[sid] = hashlib.sha256(raw.encode("utf-8")).hexdigest()
    else:
        for key in [
            "cover",
            "contenido",
            "letter_intro",
            "intro",
            "responsibility",
            "limitations",
            "closing",
            "major_interest_findings",
            "internal_control_findings",
        ]:
            value = document.get(key)
            if value is None:
                continue
            raw = json.dumps(value, ensure_ascii=False, sort_keys=True)
            hashes[key] = hashlib.sha256(raw.encode("utf-8")).hexdigest()
        findings = document.get("findings")
        if isinstance(findings, list):
            for idx, finding in enumerate(findings, start=1):
                if not isinstance(finding, dict):
                    continue
                fid = _safe_text(finding.get("id")) or str(idx)
                raw = json.dumps(finding, ensure_ascii=False, sort_keys=True)
                hashes[f"hallazgo:{fid}"] = hashlib.sha256(raw.encode("utf-8")).hexdigest()
    return hashes


def _input_hash(input_payload: dict[str, Any]) -> str:
    raw = json.dumps(input_payload or {}, ensure_ascii=False, sort_keys=True)
    return hashlib.sha256(raw.encode("utf-8")).hexdigest()


def _build_regen_diff(
    *,
    previous: dict[str, Any] | None,
    current_section_hashes: dict[str, str],
    generation_metadata: dict[str, Any],
) -> dict[str, Any]:
    if not isinstance(previous, dict):
        return {
            "has_previous": False,
            "changed_sections": sorted(current_section_hashes.keys()),
            "prompt_changed": False,
            "template_changed": False,
            "input_changed": False,
        }
    prev_meta = previous.get("generation_metadata") if isinstance(previous.get("generation_metadata"), dict) else {}
    prev_hashes = previous.get("section_hashes") if isinstance(previous.get("section_hashes"), dict) else {}
    changed = []
    all_keys = sorted(set(prev_hashes.keys()) | set(current_section_hashes.keys()))
    for key in all_keys:
        if _safe_text(prev_hashes.get(key)) != _safe_text(current_section_hashes.get(key)):
            changed.append(key)
    prompt_changed = (
        _safe_text(prev_meta.get("prompt_id")) != _safe_text(generation_metadata.get("prompt_id"))
        or _safe_text(prev_meta.get("prompt_version")) != _safe_text(generation_metadata.get("prompt_version"))
    )
    template_changed = _safe_text(prev_meta.get("template_version")) != _safe_text(generation_metadata.get("template_version"))
    input_changed = _safe_text(previous.get("input_hash")) != _input_hash(generation_metadata.get("input_payload") or {})
    return {
        "has_previous": True,
        "changed_sections": changed,
        "prompt_changed": prompt_changed,
        "template_changed": template_changed,
        "input_changed": input_changed,
    }


def _build_version_summary(diff: dict[str, Any]) -> str:
    if not isinstance(diff, dict) or not diff.get("has_previous"):
        return "Version inicial generada."
    reasons: list[str] = []
    if bool(diff.get("input_changed")):
        reasons.append("cambio de input")
    if bool(diff.get("template_changed")):
        reasons.append("cambio de plantilla")
    if bool(diff.get("prompt_changed")):
        reasons.append("cambio de prompt/modelo")
    changed_sections = diff.get("changed_sections")
    changed_count = len(changed_sections) if isinstance(changed_sections, list) else 0
    if not reasons and changed_count <= 0:
        return "Regenerado sin cambios detectados en secciones."
    reason_txt = ", ".join(reasons) if reasons else "ajuste de contenido"
    return f"Se regenero por {reason_txt}; se modificaron {changed_count} secciones."


def _register_generated_document(
    *,
    cliente_id: str,
    kind: str,
    document: dict[str, Any],
    artifacts: list[dict[str, Any]],
    generation_metadata: dict[str, Any],
) -> dict[str, Any]:
    doc_type = _document_type_for_kind(kind)
    registry = _read_registry(cliente_id)
    docs = registry.get("documents")
    if not isinstance(docs, dict):
        docs = {}
        registry["documents"] = docs

    bucket = docs.get(doc_type)
    if not isinstance(bucket, dict):
        bucket = {"document_type": doc_type, "versions": []}
        docs[doc_type] = bucket
    versions = bucket.get("versions")
    if not isinstance(versions, list):
        versions = []
        bucket["versions"] = versions

    previous_current = None
    for v in versions:
        if isinstance(v, dict) and bool(v.get("is_current")):
            previous_current = v
            break

    next_version = 1
    if versions:
        max_ver = 0
        for v in versions:
            if isinstance(v, dict):
                try:
                    max_ver = max(max_ver, int(v.get("document_version") or 0))
                except Exception:
                    continue
        next_version = max_ver + 1

    for v in versions:
        if isinstance(v, dict):
            v["is_current"] = False

    section_hashes = _section_hashes(document)
    diff = _build_regen_diff(
        previous=previous_current if isinstance(previous_current, dict) else None,
        current_section_hashes=section_hashes,
        generation_metadata=generation_metadata,
    )
    created_at = datetime.now(timezone.utc).isoformat()
    rec = {
        "document_type": doc_type,
        "document_version": next_version,
        "supersedes_version": int(previous_current.get("document_version")) if isinstance(previous_current, dict) else None,
        "is_current": True,
        "state": "draft",
        "created_at": created_at,
        "updated_at": created_at,
        "artifacts": artifacts,
        "generation_metadata": generation_metadata,
        "section_hashes": section_hashes,
        "input_hash": _input_hash(generation_metadata.get("input_payload") or {}),
        "diff_from_previous": diff,
        "summary": _build_version_summary(diff),
        "document_snapshot": document,
    }
    versions.append(rec)
    _write_registry(cliente_id, registry)
    sections = _extract_sections(document, generation_metadata.get("required_sections") if isinstance(generation_metadata.get("required_sections"), list) else [])
    _initialize_section_links_for_version(
        cliente_id=cliente_id,
        document_type=doc_type,
        document_version=next_version,
        sections=sections,
    )
    return rec


def _transition_document_state(
    *,
    cliente_id: str,
    document_type: str,
    target_state: str,
    changed_by: str,
    changed_role: str,
    reason: str = "",
) -> dict[str, Any]:
    if target_state not in DOC_STATES:
        raise HTTPException(status_code=status.HTTP_422_UNPROCESSABLE_CONTENT, detail="target_state invalido.")
    registry = _read_registry(cliente_id)
    docs = registry.get("documents")
    if not isinstance(docs, dict) or document_type not in docs:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="No existe documento para ese tipo.")
    bucket = docs.get(document_type)
    if not isinstance(bucket, dict):
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Registro documental invalido.")
    versions = bucket.get("versions")
    if not isinstance(versions, list) or not versions:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="No hay versiones documentales.")

    current = None
    for v in versions:
        if isinstance(v, dict) and bool(v.get("is_current")):
            current = v
            break
    if not isinstance(current, dict):
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="No hay version vigente.")

    current_state = _safe_text(current.get("state")) or "draft"
    if current_state == target_state:
        return current
    role_n = _normalize_role(changed_role)
    if current_state == "issued" and target_state != "issued":
        raise HTTPException(
            status_code=status.HTTP_422_UNPROCESSABLE_CONTENT,
            detail="No se permite regresar un documento emitido; debe generarse una nueva version.",
        )
    allowed_targets = _allowed_next_states(role=role_n, current_state=current_state)
    if target_state not in ALLOWED_TRANSITIONS.get(current_state, {}):
        raise HTTPException(
            status_code=status.HTTP_422_UNPROCESSABLE_CONTENT,
            detail=f"Transicion invalida: {current_state} -> {target_state}",
        )
    if target_state not in allowed_targets:
        raise HTTPException(
            status_code=status.HTTP_403_FORBIDDEN,
            detail=(
                f"Rol '{role_n}' no autorizado para transicion {current_state} -> {target_state}. "
                f"Permitido para: {', '.join(ALLOWED_TRANSITIONS.get(current_state, {}).get(target_state, []))}"
            ),
        )

    artifacts = current.get("artifacts") if isinstance(current.get("artifacts"), list) else []
    has_docx = any(_safe_text(a.get("artifact_type")) == "docx" for a in artifacts if isinstance(a, dict))
    if target_state == "approved" and not has_docx:
        raise HTTPException(
            status_code=status.HTTP_422_UNPROCESSABLE_CONTENT,
            detail="No se puede aprobar sin artefacto DOCX generado.",
        )
    if target_state == "issued" and current_state != "approved":
        raise HTTPException(
            status_code=status.HTTP_422_UNPROCESSABLE_CONTENT,
            detail="No se puede emitir si no esta en estado approved.",
        )
    current["cliente_id"] = cliente_id
    current["document_type"] = document_type
    if target_state == "approved":
        qc = _quality_check_version(current, cliente_id=cliente_id)
        failing = [c for c in qc.get("checks", []) if isinstance(c, dict) and c.get("status") != "ok"]
        if failing:
            failed_labels = ", ".join([_safe_text(c.get("code")) for c in failing])
            raise HTTPException(
                status_code=status.HTTP_422_UNPROCESSABLE_CONTENT,
                detail=f"Checklist de calidad incompleto para aprobar: {failed_labels}",
            )
        policy = _document_evidence_policy(document_type)
        if bool(policy.get("enforce_evidence_gate_on_approved", False)):
            ver = int(current.get("document_version") or 0)
            sections = _get_section_links_for_current(cliente_id, document_type, ver)
            critical_ids = _critical_hallazgo_ids(
                current.get("document_snapshot") if isinstance(current.get("document_snapshot"), dict) else {}
            )
            gate = _compute_evidence_gate(
                sections=sections,
                policy=policy,
                quality_check=qc,
                critical_hallazgo_ids=critical_ids,
            )
            if not bool(gate.get("can_approve")):
                reasons = gate.get("approve_blocking_reasons") if isinstance(gate.get("approve_blocking_reasons"), list) else []
                message = "; ".join([_safe_text(x) for x in reasons if _safe_text(x)])
                blocked_ids = [
                    _safe_text(s.get("section_id"))
                    for s in (gate.get("blocking_sections") if isinstance(gate.get("blocking_sections"), list) else [])
                    if isinstance(s, dict) and _safe_text(s.get("section_id"))
                ]
                detail = f"No se puede aprobar: {message}" if message else "No se puede aprobar por evidencia insuficiente."
                if blocked_ids:
                    detail += f" Secciones críticas sin evidencia: {', '.join(blocked_ids)}."
                raise HTTPException(status_code=status.HTTP_422_UNPROCESSABLE_CONTENT, detail=detail)
    if target_state == "issued":
        coherence = _coherence_issues_for_document(cliente_id=cliente_id, version=current, for_issue=True)
        if coherence:
            raise HTTPException(
                status_code=status.HTTP_422_UNPROCESSABLE_CONTENT,
                detail=f"No se puede emitir: {'; '.join(coherence)}",
            )
        policy = _document_evidence_policy(document_type)
        if bool(policy.get("enforce_evidence_gate_on_issued", False)):
            ver = int(current.get("document_version") or 0)
            sections = _get_section_links_for_current(cliente_id, document_type, ver)
            qc = _quality_check_version(current, cliente_id=cliente_id)
            critical_ids = _critical_hallazgo_ids(
                current.get("document_snapshot") if isinstance(current.get("document_snapshot"), dict) else {}
            )
            gate = _compute_evidence_gate(
                sections=sections,
                policy=policy,
                quality_check=qc,
                critical_hallazgo_ids=critical_ids,
            )
            if not bool(gate.get("can_issue")):
                reasons = gate.get("issue_blocking_reasons") if isinstance(gate.get("issue_blocking_reasons"), list) else []
                message = "; ".join([_safe_text(x) for x in reasons if _safe_text(x)])
                blocked_ids = [
                    _safe_text(s.get("section_id"))
                    for s in (gate.get("blocking_sections") if isinstance(gate.get("blocking_sections"), list) else [])
                    if isinstance(s, dict) and _safe_text(s.get("section_id"))
                ]
                detail = f"No se puede emitir: {message}" if message else "No se puede emitir por evidencia insuficiente."
                if blocked_ids:
                    detail += f" Secciones críticas sin evidencia: {', '.join(blocked_ids)}."
                raise HTTPException(status_code=status.HTTP_422_UNPROCESSABLE_CONTENT, detail=detail)

    transition_record = {
        "changed_by": changed_by,
        "changed_role": role_n,
        "changed_at": datetime.now(timezone.utc).isoformat(),
        "from_state": current_state,
        "to_state": target_state,
        "reason": _safe_text(reason),
    }
    state_history = current.get("state_history")
    if not isinstance(state_history, list):
        state_history = []
    state_history.append(transition_record)
    current["state_history"] = state_history[-50:]
    current["state"] = target_state
    current["updated_at"] = transition_record["changed_at"]
    current["updated_by"] = changed_by
    _write_registry(cliente_id, registry)
    return current


def _extract_section_ids(document_snapshot: dict[str, Any]) -> set[str]:
    ids: set[str] = set()
    if not isinstance(document_snapshot, dict):
        return ids
    sections = document_snapshot.get("sections")
    if isinstance(sections, list):
        for sec in sections:
            if not isinstance(sec, dict):
                continue
            sid = _safe_text(sec.get("id")) or _safe_text(sec.get("title"))
            if sid:
                ids.add(sid)
    else:
        for sid in [
            "cover",
            "contenido",
            "letter_intro",
            "intro",
            "responsibility",
            "limitations",
            "major_interest_findings",
            "internal_control_findings",
            "findings",
            "closing",
        ]:
            if sid in document_snapshot:
                ids.add(sid)
    return ids


def _extract_sections(document_snapshot: dict[str, Any], required_sections: list[str]) -> list[dict[str, Any]]:
    out: list[dict[str, Any]] = []
    req = {str(x) for x in required_sections if str(x).strip()}
    if not isinstance(document_snapshot, dict):
        return out
    if isinstance(document_snapshot.get("sections"), list):
        for sec in document_snapshot.get("sections", []):
            if not isinstance(sec, dict):
                continue
            sid = _safe_text(sec.get("id")) or _safe_text(sec.get("title"))
            if not sid:
                continue
            title = _safe_text(sec.get("title")) or sid
            out.append(
                {
                    "section_id": sid,
                    "section_title": title,
                    "is_required": sid in req,
                    "status": "pending",
                    "sources": [],
                }
            )
    else:
        fallback_sections = [
            ("cover", "Portada"),
            ("contenido", "Contenido"),
            ("letter_intro", "Bloque de carta"),
            ("intro", "Introducción"),
            ("responsibility", "Responsabilidad de la administración"),
            ("limitations", "Limitaciones"),
            ("major_interest_findings", "Aspectos de mayor interés"),
            ("internal_control_findings", "Aspectos de control interno"),
            ("closing", "Cierre"),
        ]
        for sid, title in fallback_sections:
            if sid in document_snapshot:
                out.append(
                    {
                        "section_id": sid,
                        "section_title": title,
                        "is_required": sid in req,
                        "status": "pending",
                        "sources": [],
                    }
                )
        findings = document_snapshot.get("findings")
        if not isinstance(findings, list):
            major = (
                document_snapshot.get("major_interest_findings")
                if isinstance(document_snapshot.get("major_interest_findings"), list)
                else []
            )
            control = (
                document_snapshot.get("internal_control_findings")
                if isinstance(document_snapshot.get("internal_control_findings"), list)
                else []
            )
            findings = [*major, *control]
        if isinstance(findings, list):
            for idx, finding in enumerate(findings, start=1):
                if not isinstance(finding, dict):
                    continue
                fid = _safe_text(finding.get("id")) or f"finding_{idx}"
                out.append(
                    {
                        "section_id": f"hallazgo:{fid}",
                        "section_title": _safe_text(finding.get("titulo")) or f"Hallazgo {idx}",
                        "is_required": True,
                        "status": "pending",
                        "sources": [],
                    }
                )
    return out


def _initialize_section_links_for_version(
    *,
    cliente_id: str,
    document_type: str,
    document_version: int,
    sections: list[dict[str, Any]],
) -> None:
    links = _read_links(cliente_id)
    docs = links.get("documents")
    if not isinstance(docs, dict):
        docs = {}
        links["documents"] = docs
    doc_bucket = docs.get(document_type)
    if not isinstance(doc_bucket, dict):
        doc_bucket = {"versions": {}}
        docs[document_type] = doc_bucket
    versions = doc_bucket.get("versions")
    if not isinstance(versions, dict):
        versions = {}
        doc_bucket["versions"] = versions
    versions[str(document_version)] = {"sections": sections}
    _write_links(cliente_id, links)


def _quality_check_version(version: dict[str, Any], *, cliente_id: str = "") -> dict[str, Any]:
    artifacts = version.get("artifacts") if isinstance(version.get("artifacts"), list) else []
    generation_metadata = (
        version.get("generation_metadata") if isinstance(version.get("generation_metadata"), dict) else {}
    )
    doc = version.get("document_snapshot") if isinstance(version.get("document_snapshot"), dict) else {}
    required_sections = generation_metadata.get("required_sections")
    required_sections = required_sections if isinstance(required_sections, list) else []
    section_ids = _extract_section_ids(doc)
    missing_sections = [str(s) for s in required_sections if str(s) not in section_ids]
    has_docx = any(_safe_text(a.get("artifact_type")) == "docx" for a in artifacts if isinstance(a, dict))
    doc_raw = json.dumps(doc, ensure_ascii=False).strip()
    has_pending = "[[PENDIENTE]]" in doc_raw
    has_empty_placeholder = "{{" in doc_raw or "}}" in doc_raw
    required_meta_fields = ["source", "template_mode", "template_version", "document_type", "requested_by", "generated_at"]
    missing_meta = [k for k in required_meta_fields if not _safe_text(generation_metadata.get(k))]
    checks = [
        {
            "code": "HAS_DOCX",
            "label": "Existe artefacto DOCX",
            "status": "ok" if has_docx else "blocked",
            "detail": "OK" if has_docx else "No existe artefacto DOCX.",
        },
        {
            "code": "NO_PENDING",
            "label": "Sin marcadores [[PENDIENTE]]",
            "status": "ok" if not has_pending else "blocked",
            "detail": "OK" if not has_pending else "El documento aun contiene [[PENDIENTE]].",
        },
        {
            "code": "REQUIRED_SECTIONS",
            "label": "Secciones requeridas completas",
            "status": "ok" if not missing_sections else "blocked",
            "detail": "OK" if not missing_sections else f"Faltan secciones: {', '.join(missing_sections)}",
        },
        {
            "code": "MIN_METADATA",
            "label": "Metadata minima presente",
            "status": "ok" if not missing_meta else "blocked",
            "detail": "OK" if not missing_meta else f"Faltan campos: {', '.join(missing_meta)}",
        },
        {
            "code": "EMPTY_PLACEHOLDERS",
            "label": "Sin placeholders vacios de template",
            "status": "ok" if not has_empty_placeholder else "blocked",
            "detail": "OK" if not has_empty_placeholder else "Se detectaron placeholders '{{...}}' sin resolver.",
        },
    ]
    coherence_issues = _coherence_issues_for_document(
        cliente_id=cliente_id or _safe_text(version.get("cliente_id")),
        version=version,
        for_issue=False,
    )
    checks.append(
        {
            "code": "PERIOD_COHERENCE",
            "label": "Periodo de documento coherente con perfil",
            "status": "ok" if not coherence_issues else "blocked",
            "detail": "OK" if not coherence_issues else "; ".join(coherence_issues),
        }
    )
    can_approve = all(c.get("status") == "ok" for c in checks)
    total = len(checks)
    ok_count = len([c for c in checks if c.get("status") == "ok"])
    score = int(round((ok_count / total) * 100)) if total else 0
    state = _safe_text(version.get("state")).lower()
    if state == "issued":
        semaphore = "green"
    elif can_approve and state in {"approved", "reviewed"}:
        semaphore = "green"
    elif state in {"reviewed", "approved"}:
        semaphore = "yellow"
    elif can_approve:
        semaphore = "yellow"
    else:
        semaphore = "red"
    return {"can_approve": can_approve, "checks": checks, "score": score, "semaphore": semaphore}


def _build_html_from_docx(docx_path: Path) -> str:
    try:
        from docx import Document
    except Exception as exc:
        raise HTTPException(
            status_code=status.HTTP_422_UNPROCESSABLE_CONTENT,
            detail="No se pudo emitir PDF: dependencia python-docx no disponible.",
        ) from exc
    try:
        doc = Document(str(docx_path))
    except Exception as exc:
        raise HTTPException(
            status_code=status.HTTP_422_UNPROCESSABLE_CONTENT,
            detail="No se pudo abrir el DOCX seleccionado para emisión final.",
        ) from exc
    blocks: list[str] = []
    for p in doc.paragraphs:
        text = _safe_text(p.text)
        if text:
            blocks.append(f"<p>{text}</p>")
    if not blocks:
        blocks.append("<p>Documento emitido sin contenido textual extraíble.</p>")
    body = "\n".join(blocks)
    return f"""
    <html>
      <head>
        <meta charset="utf-8" />
        <style>
          body {{ font-family: Helvetica, Arial, sans-serif; color: #0f172a; font-size: 12px; }}
          .issued {{ color: #166534; font-weight: bold; border: 2px solid #166534; display: inline-block; padding: 4px 8px; margin-bottom: 10px; }}
          p {{ margin: 4px 0; line-height: 1.35; }}
        </style>
      </head>
      <body>
        <div class="issued">EMITIDO</div>
        {body}
      </body>
    </html>
    """


def _emit_pdf_from_docx(*, docx_path: Path) -> bytes:
    html = _build_html_from_docx(docx_path)
    return _html_to_pdf_bytes(html)


def _report_requirements(cliente_id: str, user: UserContext) -> dict[str, Any]:
    dashboard = get_dashboard(cliente_id=cliente_id, user=user)
    top_areas = list(dashboard.top_areas or [])
    hallazgos = read_hallazgos(cliente_id).strip()
    generated = _generate_tasks(cliente_id)
    merged = _merge_saved_tasks(cliente_id, generated)
    gates, coverage = _quality_gates(cliente_id, merged)
    gates_map = {g.code: g.status for g in gates}

    missing_sections: list[str] = []
    if not top_areas:
        missing_sections.append("No hay areas priorizadas para el resumen ejecutivo.")
    if not hallazgos:
        missing_sections.append("No existe conclusion tecnica en hallazgos.")
    if gates_map.get("REPORT") != "ok":
        missing_sections.append("Gate REPORT debe estar en estado ok para emitir informe.")
    if coverage.total_assertions > 0 and coverage.coverage_pct <= 0:
        missing_sections.append("No existe cobertura de afirmaciones documentada para las areas del cliente.")

    can_emit_final = len(missing_sections) == 0
    can_emit_draft = bool(top_areas)
    return {
        "dashboard": dashboard,
        "top_areas": top_areas,
        "hallazgos": hallazgos,
        "gates": gates,
        "gates_map": gates_map,
        "coverage": coverage,
        "missing_sections": missing_sections,
        "can_emit_draft": can_emit_draft,
        "can_emit_final": can_emit_final,
    }


def _ensure_required_sections(cliente_id: str, user: UserContext, *, final_mode: bool) -> tuple[Any, list[Any], str]:
    report = _report_requirements(cliente_id, user)
    dashboard = report["dashboard"]
    top_areas = report["top_areas"]
    hallazgos = report["hallazgos"]
    missing_sections = report["missing_sections"]
    can_emit_draft = bool(report["can_emit_draft"])
    can_emit_final = bool(report["can_emit_final"])
    gates_map = report["gates_map"]

    if final_mode and not can_emit_final:
        raise HTTPException(
            status_code=status.HTTP_422_UNPROCESSABLE_ENTITY,
            detail={
                "message": "Faltan secciones obligatorias para generar el PDF final.",
                "errors": missing_sections,
                "gates": gates_map,
            },
        )
    if (not final_mode) and (not can_emit_draft):
        raise HTTPException(
            status_code=status.HTTP_422_UNPROCESSABLE_ENTITY,
            detail={
                "message": "No hay informacion minima para emitir borrador interno.",
                "errors": ["No hay areas priorizadas para el resumen ejecutivo."],
                "gates": gates_map,
            },
        )
    if not hallazgos:
        hallazgos = "Borrador interno sin conclusion consolidada aun. Completar hallazgos para emision final."
    return dashboard, top_areas, hallazgos


def _render_executive_html(*, dashboard: Any, top_areas: list[Any], hallazgos: str) -> str:
    rows = ""
    for item in top_areas[:6]:
        rows += (
            "<tr>"
            f"<td>{_safe_text(item.codigo)}</td>"
            f"<td>{_safe_text(item.nombre)}</td>"
            f"<td>{float(item.score_riesgo):.2f}</td>"
            f"<td>{_safe_text(item.prioridad).upper()}</td>"
            "</tr>"
        )

    return f"""
    <html>
      <head>
        <meta charset="utf-8" />
        <style>
          body {{ font-family: Helvetica, Arial, sans-serif; color: #0f172a; font-size: 12px; }}
          .header {{ background: linear-gradient(135deg, #041627 0%, #1a2b3c 100%); color: #fff; padding: 22px; border-radius: 10px; }}
          .small {{ color: #64748b; font-size: 11px; text-transform: uppercase; letter-spacing: 0.08em; }}
          .title {{ font-size: 27px; margin: 6px 0 0 0; }}
          .grid {{ margin-top: 16px; }}
          .card {{ border: 1px solid #cbd5e1; border-radius: 8px; padding: 10px; margin-bottom: 8px; }}
          table {{ width: 100%; border-collapse: collapse; margin-top: 10px; }}
          th, td {{ border: 1px solid #cbd5e1; padding: 6px; text-align: left; }}
          th {{ background: #f1f5f9; }}
          .footer {{ margin-top: 24px; color: #64748b; font-size: 10px; }}
          pre {{ white-space: pre-wrap; line-height: 1.35; font-size: 11px; background: #f8fafc; border: 1px solid #e2e8f0; padding: 10px; border-radius: 8px; }}
        </style>
      </head>
      <body>
        <div class="header">
          <div class="small">Socio AI · Reporte Ejecutivo</div>
          <div class="title">{_safe_text(dashboard.nombre_cliente)}</div>
          <div>Periodo: {_safe_text(dashboard.periodo)} · Riesgo Global: {_safe_text(dashboard.riesgo_global).upper()}</div>
        </div>

        <div class="grid">
          <div class="card"><b>Materialidad de planeacion:</b> ${float(dashboard.materialidad_global):,.2f}</div>
          <div class="card"><b>Materialidad de ejecucion:</b> ${float(dashboard.materialidad_ejecucion):,.2f}</div>
          <div class="card"><b>Umbral trivial:</b> ${float(dashboard.umbral_trivial):,.2f}</div>
        </div>

        <h3>Top areas de riesgo</h3>
        <table>
          <thead>
            <tr><th>Area</th><th>Nombre</th><th>Score</th><th>Prioridad</th></tr>
          </thead>
          <tbody>{rows}</tbody>
        </table>

        <h3>Conclusion tecnica</h3>
        <pre>{hallazgos}</pre>

        <div class="footer">
          Generado: {datetime.now(timezone.utc).isoformat()} · Socio AI Executive Report
        </div>
      </body>
    </html>
    """


def _html_to_pdf_bytes(html: str) -> bytes:
    try:
        from xhtml2pdf import pisa
    except Exception as exc:
        raise HTTPException(
            status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
            detail="Dependencia xhtml2pdf no disponible en backend.",
        ) from exc

    output = BytesIO()
    result = pisa.CreatePDF(src=html, dest=output)
    if result.err:
        raise HTTPException(status_code=status.HTTP_500_INTERNAL_SERVER_ERROR, detail="No se pudo renderizar PDF.")
    return output.getvalue()


def _build_memo_fallback(*, dashboard: Any, top_areas: list[Any], hallazgos: str) -> str:
    areas = ", ".join([f"{x.codigo} ({x.nombre})" for x in top_areas[:3]]) if top_areas else "sin areas destacadas"
    short_h = hallazgos.strip()
    if len(short_h) > 900:
        short_h = short_h[:900] + "..."
    return (
        f"Memo Ejecutivo - {dashboard.nombre_cliente}\n\n"
        f"Riesgo global: {dashboard.riesgo_global}.\n"
        f"Materialidad de planeacion: {dashboard.materialidad_global:,.2f}.\n"
        f"Areas prioritarias: {areas}.\n\n"
        "Recomendacion:\n"
        "1) Cerrar procedimientos pendientes en areas altas.\n"
        "2) Documentar evidencia de soporte y conclusion por area.\n"
        "3) Validar consistencia entre TB, Mayor y revelaciones.\n\n"
        f"Contexto de hallazgos:\n{short_h}"
    )


@router.get("/{cliente_id}/executive-pdf", response_model=ApiResponse)
def get_executive_pdf(
    cliente_id: str,
    mode: str = Query("draft", pattern="^(draft|final)$"),
    user: UserContext = Depends(get_current_user),
) -> ApiResponse:
    authorize_cliente_access(cliente_id, user)

    final_mode = mode == "final"
    dashboard, top_areas, hallazgos = _ensure_required_sections(cliente_id, user, final_mode=final_mode)
    html = _render_executive_html(dashboard=dashboard, top_areas=top_areas, hallazgos=hallazgos)
    pdf_bytes = _html_to_pdf_bytes(html)

    EXPORTS_DIR.mkdir(parents=True, exist_ok=True)
    timestamp = datetime.now(timezone.utc).strftime("%Y%m%dT%H%M%SZ")
    report_name = f"{cliente_id}_executive_summary_{timestamp}.pdf"
    out_path = EXPORTS_DIR / report_name
    out_path.write_bytes(pdf_bytes)

    file_hash = hashlib.sha256(pdf_bytes).hexdigest()
    _append_history(
        cliente_id,
        {
            "kind": "executive_pdf",
            "report_name": report_name,
            "generated_at": datetime.now(timezone.utc).isoformat(),
            "path": str(out_path),
            "file_hash": file_hash,
            "size_bytes": len(pdf_bytes),
            "status": "success",
            "origin": "motor_html_pdf_final" if final_mode else "motor_html_pdf_draft",
        },
    )
    payload = PdfSummaryResponse(
        cliente_id=cliente_id,
        report_name=report_name,
        generated_at=datetime.now(timezone.utc),
        path=str(out_path),
        file_hash=file_hash,
        size_bytes=len(pdf_bytes),
    )
    return ApiResponse(data=payload.model_dump())


@router.get("/{cliente_id}/executive-pdf/file")
def get_executive_pdf_file(
    cliente_id: str,
    path: str = Query(...),
    user: UserContext = Depends(get_current_user),
) -> FileResponse:
    authorize_cliente_access(cliente_id, user)
    target = Path(path)
    if not target.is_absolute():
        target = (ROOT / target).resolve()
    else:
        target = target.resolve()

    exports_resolved = EXPORTS_DIR.resolve()
    if not str(target).startswith(str(exports_resolved)):
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Ruta de archivo invalida.")
    if not target.exists():
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Archivo no encontrado.")
    return FileResponse(path=target, filename=target.name, media_type="application/pdf")


@router.post("/{cliente_id}/memo", response_model=ApiResponse)
def post_executive_memo(cliente_id: str, user: UserContext = Depends(get_current_user)) -> ApiResponse:
    authorize_cliente_access(cliente_id, user)
    dashboard, top_areas, hallazgos = _ensure_required_sections(cliente_id, user, final_mode=False)
    query = (
        "Genera un memo ejecutivo de auditoria en espanol para socio firmante. "
        "Incluye: riesgo global, materialidad, top areas, recomendaciones accionables y siguiente paso."
    )
    rag = generate_chat_response(cliente_id, query)
    memo = str(rag.get("answer") or "").strip()
    provider = str(rag.get("provider") or "llm")
    if not memo:
        memo = _build_memo_fallback(dashboard=dashboard, top_areas=top_areas, hallazgos=hallazgos)

    write_memo(cliente_id, memo)
    marker = f"## Memo Ejecutivo ({datetime.now(timezone.utc).date().isoformat()})"
    append_hallazgo(cliente_id, f"{marker}\n\n{memo}")
    _append_history(
        cliente_id,
        {
            "kind": "executive_memo",
            "report_name": f"{cliente_id}_memo_{datetime.now(timezone.utc).strftime('%Y%m%dT%H%M%SZ')}",
            "generated_at": datetime.now(timezone.utc).isoformat(),
            "path": "",
            "file_hash": hashlib.sha256(memo.encode("utf-8")).hexdigest(),
            "size_bytes": len(memo.encode("utf-8")),
            "status": "success",
            "origin": f"{provider}_rag" if str(rag.get("answer") or "").strip() else "fallback",
        },
    )

    payload = ReportMemoResponse(
        cliente_id=cliente_id,
        memo=memo,
        generated_at=datetime.now(timezone.utc),
        source=f"{provider}_rag" if str(rag.get("answer") or "").strip() else "fallback",
    )
    return ApiResponse(data=payload.model_dump())


@router.get("/{cliente_id}/memo", response_model=ApiResponse)
def get_executive_memo(cliente_id: str, user: UserContext = Depends(get_current_user)) -> ApiResponse:
    authorize_cliente_access(cliente_id, user)
    memo = read_memo(cliente_id)
    payload = ReportMemoResponse(
        cliente_id=cliente_id,
        memo=memo,
        generated_at=datetime.now(timezone.utc),
        source="saved" if memo else "empty",
    )
    return ApiResponse(data=payload.model_dump())


@router.post("/{cliente_id}/carta-control-interno", response_model=ApiResponse)
def post_internal_control_letter(
    cliente_id: str,
    payload: InternalControlLetterRequest,
    user: UserContext = Depends(get_current_user),
) -> ApiResponse:
    authorize_cliente_access(cliente_id, user)
    normalized_recipient = (payload.recipient or "").strip() or "Gerencia General"
    generated = generate_internal_control_letter(
        cliente_id,
        recipient=normalized_recipient,
        include_management_response=payload.include_management_response,
        max_findings=payload.max_findings,
        requested_by=user.sub,
    )
    content = str(generated.get("content") or "")
    out_path = str(generated.get("path") or "")
    document = generated.get("document") if isinstance(generated.get("document"), dict) else {}
    artifacts = generated.get("artifacts") if isinstance(generated.get("artifacts"), list) else []
    generation_metadata = (
        generated.get("generation_metadata") if isinstance(generated.get("generation_metadata"), dict) else {}
    )
    source = str(generation_metadata.get("source") or "fallback")
    findings_count = int(generated.get("findings_count") or 0)
    version_rec = _register_generated_document(
        cliente_id=cliente_id,
        kind="internal_control_letter",
        document=document,
        artifacts=artifacts,
        generation_metadata=generation_metadata,
    )

    _append_history(
        cliente_id,
        {
            "kind": "internal_control_letter",
            "report_name": Path(out_path).name if out_path else f"{cliente_id}_carta_control_interno",
            "generated_at": datetime.now(timezone.utc).isoformat(),
            "path": out_path,
            "file_hash": hashlib.sha256(content.encode("utf-8")).hexdigest(),
            "size_bytes": len(content.encode("utf-8")),
            "status": "success",
            "origin": source,
            "findings_count": findings_count,
            "document_version": version_rec.get("document_version"),
            "supersedes_version": version_rec.get("supersedes_version"),
            "is_current": version_rec.get("is_current"),
            "state": version_rec.get("state"),
            "diff_from_previous": version_rec.get("diff_from_previous") or {},
            "generation_metadata": {
                "source": source,
                "provider": str(generation_metadata.get("provider") or ""),
                "model": str(generation_metadata.get("model") or ""),
                "prompt_id": str(generation_metadata.get("prompt_id") or ""),
                "prompt_version": str(generation_metadata.get("prompt_version") or ""),
                "document_type": str(generation_metadata.get("document_type") or ""),
                "template_mode": str(generation_metadata.get("template_mode") or ""),
                "template_version": str(generation_metadata.get("template_version") or ""),
                "placeholders_supported": generation_metadata.get("placeholders_supported") or [],
                "required_sections": generation_metadata.get("required_sections") or [],
                "optional_sections": generation_metadata.get("optional_sections") or [],
                "input_payload": generation_metadata.get("input_payload") or {},
                "requested_by": str(generation_metadata.get("requested_by") or user.sub),
                "generated_at": str(generation_metadata.get("generated_at") or datetime.now(timezone.utc).isoformat()),
            },
            "artifacts": artifacts,
        },
    )

    response = InternalControlLetterResponse(
        cliente_id=cliente_id,
        generated_at=datetime.now(timezone.utc),
        recipient=normalized_recipient,
        findings_count=findings_count,
        source=source,
        document_version=int(version_rec.get("document_version") or 1),
        supersedes_version=version_rec.get("supersedes_version"),
        is_current=bool(version_rec.get("is_current", True)),
        state=str(version_rec.get("state") or "draft"),
        diff_from_previous=version_rec.get("diff_from_previous") or {},
        document=document,
        generation_metadata=generation_metadata,
        artifacts=artifacts,
        content=content,
        path=out_path,
    )
    return ApiResponse(data=response.model_dump())


@router.post("/{cliente_id}/niif-pymes-borrador", response_model=ApiResponse)
def post_niif_pymes_draft(
    cliente_id: str,
    payload: NIIFPymesDraftRequest,
    user: UserContext = Depends(get_current_user),
) -> ApiResponse:
    authorize_cliente_access(cliente_id, user)
    if payload.early_adoption and payload.ifrs_for_smes_version == "2015":
        raise HTTPException(
            status_code=status.HTTP_422_UNPROCESSABLE_CONTENT,
            detail="early_adoption=true no aplica para NIIF para PYMES 2015.",
        )
    generated = generate_niif_pymes_draft(
        cliente_id,
        ifrs_for_smes_version=payload.ifrs_for_smes_version,
        early_adoption=payload.early_adoption,
        include_policy_section=payload.include_policy_section,
        requested_by=user.sub,
    )
    content = str(generated.get("content") or "")
    out_path = str(generated.get("path") or "")
    document = generated.get("document") if isinstance(generated.get("document"), dict) else {}
    artifacts = generated.get("artifacts") if isinstance(generated.get("artifacts"), list) else []
    generation_metadata = (
        generated.get("generation_metadata") if isinstance(generated.get("generation_metadata"), dict) else {}
    )
    source = str(generation_metadata.get("source") or "fallback")
    version_rec = _register_generated_document(
        cliente_id=cliente_id,
        kind="niif_pymes_draft",
        document=document,
        artifacts=artifacts,
        generation_metadata=generation_metadata,
    )

    _append_history(
        cliente_id,
        {
            "kind": "niif_pymes_draft",
            "report_name": Path(out_path).name if out_path else f"{cliente_id}_niif_pymes_borrador",
            "generated_at": datetime.now(timezone.utc).isoformat(),
            "path": out_path,
            "file_hash": hashlib.sha256(content.encode("utf-8")).hexdigest(),
            "size_bytes": len(content.encode("utf-8")),
            "status": "success",
            "origin": source,
            "ifrs_for_smes_version": payload.ifrs_for_smes_version,
            "early_adoption": payload.early_adoption,
            "document_version": version_rec.get("document_version"),
            "supersedes_version": version_rec.get("supersedes_version"),
            "is_current": version_rec.get("is_current"),
            "state": version_rec.get("state"),
            "diff_from_previous": version_rec.get("diff_from_previous") or {},
            "generation_metadata": {
                "source": source,
                "provider": str(generation_metadata.get("provider") or ""),
                "model": str(generation_metadata.get("model") or ""),
                "prompt_id": str(generation_metadata.get("prompt_id") or ""),
                "prompt_version": str(generation_metadata.get("prompt_version") or ""),
                "document_type": str(generation_metadata.get("document_type") or ""),
                "template_mode": str(generation_metadata.get("template_mode") or ""),
                "template_version": str(generation_metadata.get("template_version") or ""),
                "placeholders_supported": generation_metadata.get("placeholders_supported") or [],
                "required_sections": generation_metadata.get("required_sections") or [],
                "optional_sections": generation_metadata.get("optional_sections") or [],
                "input_payload": generation_metadata.get("input_payload") or {},
                "requested_by": str(generation_metadata.get("requested_by") or user.sub),
                "generated_at": str(generation_metadata.get("generated_at") or datetime.now(timezone.utc).isoformat()),
            },
            "artifacts": artifacts,
        },
    )

    response = NIIFPymesDraftResponse(
        cliente_id=cliente_id,
        generated_at=datetime.now(timezone.utc),
        period_end=str(generated.get("period_end") or ""),
        ifrs_for_smes_version=str(generated.get("ifrs_for_smes_version") or payload.ifrs_for_smes_version),
        early_adoption=bool(generated.get("early_adoption")),
        source=source,
        document_version=int(version_rec.get("document_version") or 1),
        supersedes_version=version_rec.get("supersedes_version"),
        is_current=bool(version_rec.get("is_current", True)),
        state=str(version_rec.get("state") or "draft"),
        diff_from_previous=version_rec.get("diff_from_previous") or {},
        document=document,
        generation_metadata=generation_metadata,
        artifacts=artifacts,
        content=content,
        path=out_path,
    )
    return ApiResponse(data=response.model_dump())


@router.get("/{cliente_id}/historial", response_model=ApiResponse)
def get_report_history(cliente_id: str, user: UserContext = Depends(get_current_user)) -> ApiResponse:
    authorize_cliente_access(cliente_id, user)
    generated = _generate_tasks(cliente_id)
    merged = _merge_saved_tasks(cliente_id, generated)
    gates, coverage = _quality_gates(cliente_id, merged)
    gate_map = {g.code: g.status for g in gates}
    history = _read_history(cliente_id)
    return ApiResponse(
        data={
            "cliente_id": cliente_id,
            "gates": [g.model_dump() for g in gates],
            "gate_status": gate_map,
            "coverage_summary": coverage.model_dump(),
            "items": history,
        }
    )


@router.get("/{cliente_id}/status", response_model=ApiResponse)
def get_report_status(cliente_id: str, user: UserContext = Depends(get_current_user)) -> ApiResponse:
    authorize_cliente_access(cliente_id, user)
    report = _report_requirements(cliente_id, user)
    return ApiResponse(
        data={
            "cliente_id": cliente_id,
            "gates": [g.model_dump() for g in report.get("gates", [])],
            "missing_sections": report.get("missing_sections", []),
            "can_emit_draft": bool(report.get("can_emit_draft")),
            "can_emit_final": bool(report.get("can_emit_final")),
            "coverage_summary": report.get("coverage").model_dump() if report.get("coverage") else {},
        }
    )


@router.get("/{cliente_id}/documentos/{document_type}/versiones", response_model=ApiResponse)
def get_document_versions(
    cliente_id: str,
    document_type: str,
    user: UserContext = Depends(get_current_user),
) -> ApiResponse:
    authorize_cliente_access(cliente_id, user)
    registry = _read_registry(cliente_id)
    docs = registry.get("documents") if isinstance(registry.get("documents"), dict) else {}
    bucket = docs.get(document_type) if isinstance(docs, dict) else None
    versions = bucket.get("versions") if isinstance(bucket, dict) and isinstance(bucket.get("versions"), list) else []
    return ApiResponse(
        data={
            "cliente_id": cliente_id,
            "document_type": document_type,
            "versions": versions,
            "current": next((v for v in versions if isinstance(v, dict) and bool(v.get("is_current"))), None),
        }
    )


def _get_current_version_record(cliente_id: str, document_type: str) -> dict[str, Any]:
    registry = _read_registry(cliente_id)
    docs = registry.get("documents") if isinstance(registry.get("documents"), dict) else {}
    bucket = docs.get(document_type) if isinstance(docs, dict) else None
    versions = bucket.get("versions") if isinstance(bucket, dict) and isinstance(bucket.get("versions"), list) else []
    current = next((v for v in versions if isinstance(v, dict) and bool(v.get("is_current"))), None)
    if not isinstance(current, dict):
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="No hay version vigente para ese documento.")
    return current


def _get_section_links_for_current(cliente_id: str, document_type: str, current_version: int) -> list[dict[str, Any]]:
    links = _read_links(cliente_id)
    docs = links.get("documents") if isinstance(links.get("documents"), dict) else {}
    doc_bucket = docs.get(document_type) if isinstance(docs, dict) else None
    versions = doc_bucket.get("versions") if isinstance(doc_bucket, dict) and isinstance(doc_bucket.get("versions"), dict) else {}
    version_bucket = versions.get(str(current_version)) if isinstance(versions, dict) else None
    sections = version_bucket.get("sections") if isinstance(version_bucket, dict) and isinstance(version_bucket.get("sections"), list) else []
    return [s for s in sections if isinstance(s, dict)]


def _section_content_from_snapshot(snapshot: dict[str, Any], section_id: str) -> Any:
    if not isinstance(snapshot, dict):
        return ""
    if section_id.startswith("hallazgo:") or section_id.startswith("finding:"):
        target = section_id.split(":", 1)[1]
        findings = snapshot.get("findings")
        if not isinstance(findings, list):
            major = snapshot.get("major_interest_findings") if isinstance(snapshot.get("major_interest_findings"), list) else []
            control = (
                snapshot.get("internal_control_findings")
                if isinstance(snapshot.get("internal_control_findings"), list)
                else []
            )
            findings = [*major, *control]
        if isinstance(findings, list):
            for f in findings:
                if not isinstance(f, dict):
                    continue
                fid = _safe_text(f.get("id"))
                if fid == target:
                    return f
        return {}
    sections = snapshot.get("sections")
    if isinstance(sections, list):
        for sec in sections:
            if not isinstance(sec, dict):
                continue
            sid = _safe_text(sec.get("id")) or _safe_text(sec.get("title"))
            if sid == section_id:
                return sec
    return snapshot.get(section_id, "")


def _refresh_section_status(section: dict[str, Any]) -> dict[str, Any]:
    sources = section.get("sources") if isinstance(section.get("sources"), list) else []
    has_sources = len([s for s in sources if isinstance(s, dict)]) > 0
    is_required = bool(section.get("is_required"))
    if has_sources:
        status = "supported"
    elif is_required:
        status = "missing_required_support"
    else:
        status = "pending"
    section["status"] = status
    return section


def _validate_link_source_exists(
    *,
    cliente_id: str,
    source_type: str,
    source_id: str,
    document_snapshot: dict[str, Any],
) -> tuple[bool, str]:
    if source_type in {"workpaper", "cedula"}:
        tasks = read_workpapers(cliente_id)
        task_ids = {_safe_text(t.get("id")) for t in tasks if isinstance(t, dict)}
        if source_id not in task_ids:
            return False, f"{source_type} '{source_id}' no existe en papeles_trabajo."
        return True, ""
    if source_type in {"trial_balance", "balance"}:
        cdir = REPORTS_META_DIR / cliente_id
        if not (cdir / "tb.xlsx").exists():
            return False, "No existe TB cargado para validar source_id de balance."
        return True, ""
    if source_type == "hallazgo":
        hallazgos_md = read_hallazgos(cliente_id)
        if source_id not in hallazgos_md:
            return False, f"hallazgo '{source_id}' no existe en hallazgos.md."
        return True, ""
    if source_type == "management_response":
        findings = document_snapshot.get("findings") if isinstance(document_snapshot.get("findings"), list) else []
        for f in findings:
            if not isinstance(f, dict):
                continue
            fid = _safe_text(f.get("id"))
            response = _safe_text(f.get("comentarios_administracion")) or _safe_text(f.get("respuesta_gerencia"))
            if fid == source_id and not _is_pending_marker(response):
                return True, ""
        return False, "management_response no encontrada o pendiente."
    return True, ""


def _compute_evidence_gate(
    *,
    sections: list[dict[str, Any]],
    policy: dict[str, Any],
    quality_check: dict[str, Any] | None = None,
    critical_hallazgo_ids: set[str] | None = None,
) -> dict[str, Any]:
    critical_sections = [
        _safe_text(x) for x in (policy.get("critical_sections") if isinstance(policy.get("critical_sections"), list) else []) if _safe_text(x)
    ]
    minimum_required = float(policy.get("minimum_coverage_percent") or 0.0)
    normalized: list[dict[str, Any]] = []
    blocking_sections: list[dict[str, Any]] = []
    blocking_keys: set[str] = set()
    supported = 0

    for raw in sections:
        section = _refresh_section_status(dict(raw))
        section_id = _safe_text(section.get("section_id"))
        sources = section.get("sources") if isinstance(section.get("sources"), list) else []
        linked_support_count = len([s for s in sources if isinstance(s, dict)])
        required_support_count = 1 if bool(section.get("is_required")) else 0
        coverage_percent = 100 if linked_support_count >= required_support_count else 0
        is_critical = any(_section_matches_critical(section_id, crit) for crit in critical_sections)
        status = _safe_text(section.get("status"))
        blocking_reason = ""
        if is_critical and status in {"missing_required_support", "pending"}:
            blocking_reason = "missing_required_support"
            key = f"{section_id}:{blocking_reason}"
            if key not in blocking_keys:
                blocking_keys.add(key)
                blocking_sections.append(
                    {
                        "section_id": section_id,
                        "section_title": _safe_text(section.get("section_title")),
                        "reason": blocking_reason,
                    }
                )
        if section_id.startswith("hallazgo:") and critical_hallazgo_ids:
            fid = section_id.split(":", 1)[1]
            if fid in critical_hallazgo_ids and status != "supported":
                blocking_reason = "critical_finding_without_evidence"
                key = f"{section_id}:{blocking_reason}"
                if key not in blocking_keys:
                    blocking_keys.add(key)
                    blocking_sections.append(
                        {
                            "section_id": section_id,
                            "section_title": _safe_text(section.get("section_title")),
                            "reason": blocking_reason,
                        }
                    )
        if status == "supported":
            supported += 1
        section["is_critical"] = is_critical
        section["required_support_count"] = required_support_count
        section["linked_support_count"] = linked_support_count
        section["coverage_percent"] = coverage_percent
        section["blocking_reason"] = blocking_reason
        normalized.append(section)

    total = len(normalized)
    coverage_percent = round((supported / total) * 100, 2) if total else 0.0
    has_missing_required = any(_safe_text(s.get("status")) == "missing_required_support" for s in normalized)
    coverage_below_minimum = coverage_percent < minimum_required

    approve_reasons: list[str] = []
    issue_reasons: list[str] = []
    if blocking_sections:
        approve_reasons.append(f"Faltan soportes en {len(blocking_sections)} secciones críticas.")
        issue_reasons.append(f"Faltan soportes en {len(blocking_sections)} secciones críticas.")
    if coverage_below_minimum:
        msg = f"Cobertura {coverage_percent:.2f}%, mínimo requerido {minimum_required:.2f}%."
        approve_reasons.append(msg)
        issue_reasons.append(msg)
    if has_missing_required:
        issue_reasons.append("Existen secciones con missing_required_support.")

    qc_can_approve = bool((quality_check or {}).get("can_approve", True))
    if not qc_can_approve:
        issue_reasons.append("Checklist documental no está en verde.")

    can_approve = not (blocking_sections or coverage_below_minimum)
    can_issue = not (blocking_sections or coverage_below_minimum or has_missing_required or not qc_can_approve)
    return {
        "policy": policy,
        "sections": normalized,
        "coverage_percent": coverage_percent,
        "minimum_required": minimum_required,
        "missing_required_support_count": len([s for s in normalized if _safe_text(s.get("status")) == "missing_required_support"]),
        "blocking_sections": blocking_sections,
        "can_approve": can_approve,
        "can_issue": can_issue,
        "approve_blocking_reasons": approve_reasons,
        "issue_blocking_reasons": issue_reasons,
    }


def _evidence_gate_for_current_version(cliente_id: str, document_type: str) -> dict[str, Any]:
    current = _get_current_version_record(cliente_id, document_type)
    document_version = int(current.get("document_version") or 0)
    sections = _get_section_links_for_current(cliente_id, document_type, document_version)
    policy = _document_evidence_policy(document_type)
    qc = _quality_check_version(current, cliente_id=cliente_id)
    critical_ids = _critical_hallazgo_ids(current.get("document_snapshot") if isinstance(current.get("document_snapshot"), dict) else {})
    gate = _compute_evidence_gate(
        sections=sections,
        policy=policy,
        quality_check=qc,
        critical_hallazgo_ids=critical_ids,
    )
    gate.update(
        {
            "cliente_id": cliente_id,
            "document_type": document_type,
            "document_version": document_version,
            "state": _safe_text(current.get("state")) or "draft",
            "enforce_on_approved": bool(policy.get("enforce_evidence_gate_on_approved", False)),
            "enforce_on_issued": bool(policy.get("enforce_evidence_gate_on_issued", False)),
        }
    )
    return gate


@router.get("/{cliente_id}/documentos/{document_type}/secciones", response_model=ApiResponse)
def get_document_sections(
    cliente_id: str,
    document_type: str,
    user: UserContext = Depends(get_current_user),
) -> ApiResponse:
    authorize_cliente_access(cliente_id, user)
    current = _get_current_version_record(cliente_id, document_type)
    ver = int(current.get("document_version") or 0)
    sections = _get_section_links_for_current(cliente_id, document_type, ver)
    policy = _document_evidence_policy(document_type)
    critical_ids = _critical_hallazgo_ids(current.get("document_snapshot") if isinstance(current.get("document_snapshot"), dict) else {})
    gate = _compute_evidence_gate(
        sections=sections,
        policy=policy,
        quality_check=_quality_check_version(current, cliente_id=cliente_id),
        critical_hallazgo_ids=critical_ids,
    )
    normalized = gate.get("sections") if isinstance(gate.get("sections"), list) else []
    missing_required = len([s for s in normalized if s.get("status") == "missing_required_support"])
    return ApiResponse(
        data={
            "cliente_id": cliente_id,
            "document_type": document_type,
            "document_version": ver,
            "sections": normalized,
            "coverage": {
                "total_sections": len(normalized),
                "supported_sections": len([s for s in normalized if s.get("status") == "supported"]),
                "missing_required": missing_required,
                "coverage_percent": gate.get("coverage_percent"),
            },
        }
    )


@router.get("/{cliente_id}/documentos/{document_type}/secciones/{section_id}/evidencia", response_model=ApiResponse)
def get_document_section_evidence(
    cliente_id: str,
    document_type: str,
    section_id: str,
    user: UserContext = Depends(get_current_user),
) -> ApiResponse:
    authorize_cliente_access(cliente_id, user)
    current = _get_current_version_record(cliente_id, document_type)
    snapshot = current.get("document_snapshot") if isinstance(current.get("document_snapshot"), dict) else {}
    ver = int(current.get("document_version") or 0)
    sections = _get_section_links_for_current(cliente_id, document_type, ver)
    section = next((s for s in sections if _safe_text(s.get("section_id")) == section_id), None)
    if not isinstance(section, dict):
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Sección no encontrada.")
    snapshot = current.get("document_snapshot") if isinstance(current.get("document_snapshot"), dict) else {}
    policy = _document_evidence_policy(document_type)
    critical_ids = _critical_hallazgo_ids(snapshot if isinstance(snapshot, dict) else {})
    gate = _compute_evidence_gate(
        sections=sections,
        policy=policy,
        quality_check=_quality_check_version(current, cliente_id=cliente_id),
        critical_hallazgo_ids=critical_ids,
    )
    gate_sections = gate.get("sections") if isinstance(gate.get("sections"), list) else []
    enriched = next(
        (s for s in gate_sections if isinstance(s, dict) and _safe_text(s.get("section_id")) == section_id),
        _refresh_section_status(dict(section)),
    )
    return ApiResponse(
        data={
            "cliente_id": cliente_id,
            "document_type": document_type,
            "document_version": ver,
            "section": enriched,
            "section_content": _section_content_from_snapshot(snapshot, section_id),
            "artifacts": current.get("artifacts") if isinstance(current.get("artifacts"), list) else [],
            "state_history": current.get("state_history") if isinstance(current.get("state_history"), list) else [],
        }
    )


@router.post("/{cliente_id}/documentos/{document_type}/secciones/{section_id}/evidencia", response_model=ApiResponse)
def post_document_section_evidence(
    cliente_id: str,
    document_type: str,
    section_id: str,
    payload: dict[str, Any],
    user: UserContext = Depends(get_current_user),
) -> ApiResponse:
    authorize_cliente_access(cliente_id, user)
    source_type = _safe_text(payload.get("source_type")).lower()
    source_id = _safe_text(payload.get("source_id"))
    reference = _safe_text(payload.get("reference"))
    label = _safe_text(payload.get("label"))
    is_required = bool(payload.get("is_required", False))
    if not source_type or not source_id or not label:
        raise HTTPException(
            status_code=status.HTTP_422_UNPROCESSABLE_CONTENT,
            detail="source_type, source_id y label son obligatorios.",
        )
    allowed_source_types = {
        "trial_balance",
        "balance",
        "workpaper",
        "hallazgo",
        "management_response",
        "adjustment",
        "supporting_document",
        "cedula",
    }
    if source_type not in allowed_source_types:
        raise HTTPException(
            status_code=status.HTTP_422_UNPROCESSABLE_CONTENT,
            detail=f"source_type inválido. Usa uno de: {', '.join(sorted(allowed_source_types))}.",
        )

    current = _get_current_version_record(cliente_id, document_type)
    snapshot = current.get("document_snapshot") if isinstance(current.get("document_snapshot"), dict) else {}
    ver = int(current.get("document_version") or 0)
    links = _read_links(cliente_id)
    docs = links.get("documents") if isinstance(links.get("documents"), dict) else {}
    doc_bucket = docs.get(document_type) if isinstance(docs, dict) else None
    versions = doc_bucket.get("versions") if isinstance(doc_bucket, dict) and isinstance(doc_bucket.get("versions"), dict) else {}
    version_bucket = versions.get(str(ver)) if isinstance(versions, dict) else None
    sections = version_bucket.get("sections") if isinstance(version_bucket, dict) and isinstance(version_bucket.get("sections"), list) else []
    target = next((s for s in sections if isinstance(s, dict) and _safe_text(s.get("section_id")) == section_id), None)
    if not isinstance(target, dict):
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Sección no encontrada para vincular evidencia.")
    source_ok, source_err = _validate_link_source_exists(
        cliente_id=cliente_id,
        source_type=source_type,
        source_id=source_id,
        document_snapshot=snapshot if isinstance(snapshot, dict) else {},
    )
    if not source_ok:
        raise HTTPException(status_code=status.HTTP_422_UNPROCESSABLE_CONTENT, detail=source_err)

    sources = target.get("sources")
    if not isinstance(sources, list):
        sources = []
        target["sources"] = sources
    composite = f"{source_type}|{source_id}|{reference}"
    existing = None
    for src in sources:
        if not isinstance(src, dict):
            continue
        if f"{_safe_text(src.get('source_type')).lower()}|{_safe_text(src.get('source_id'))}|{_safe_text(src.get('reference'))}" == composite:
            existing = src
            break
    source_payload = {
        "source_type": source_type,
        "source_id": source_id,
        "reference": reference,
        "label": label,
        "linked_by": user.sub,
        "linked_at": datetime.now(timezone.utc).isoformat(),
        "mode": "manual",
        "validated": source_ok,
    }
    if isinstance(existing, dict):
        existing.update(source_payload)
    else:
        sources.append(source_payload)
    target["is_required"] = bool(target.get("is_required")) or is_required
    _refresh_section_status(target)
    _write_links(cliente_id, links)

    _append_history(
        cliente_id,
        {
            "kind": "section_evidence_linked",
            "generated_at": datetime.now(timezone.utc).isoformat(),
            "document_type": document_type,
            "document_version": ver,
            "section_id": section_id,
            "source_type": source_type,
            "source_id": source_id,
            "status": "success",
            "changed_by": user.sub,
        },
    )
    return ApiResponse(
        data={
            "cliente_id": cliente_id,
            "document_type": document_type,
            "document_version": ver,
            "section": target,
        }
    )


@router.get("/{cliente_id}/documentos/{document_type}/acciones", response_model=ApiResponse)
def get_document_allowed_actions(
    cliente_id: str,
    document_type: str,
    user: UserContext = Depends(get_current_user),
) -> ApiResponse:
    authorize_cliente_access(cliente_id, user)
    registry = _read_registry(cliente_id)
    docs = registry.get("documents") if isinstance(registry.get("documents"), dict) else {}
    bucket = docs.get(document_type) if isinstance(docs, dict) else None
    versions = bucket.get("versions") if isinstance(bucket, dict) and isinstance(bucket.get("versions"), list) else []
    current = next((v for v in versions if isinstance(v, dict) and bool(v.get("is_current"))), None)
    current_state = _safe_text(current.get("state")) if isinstance(current, dict) else "draft"
    role_n = _normalize_role(user.role)
    allowed_next = _allowed_next_states(role=role_n, current_state=current_state)
    return ApiResponse(
        data={
            "cliente_id": cliente_id,
            "document_type": document_type,
            "role": role_n,
            "current_state": current_state,
            "permissions": [
                "document.read",
                "document.generate",
                *[f"document.transition.{x}" for x in allowed_next],
            ],
            "allowed_next_states": allowed_next,
        }
    )


@router.get("/{cliente_id}/documentos/{document_type}/quality-check", response_model=ApiResponse)
def get_document_quality_check(
    cliente_id: str,
    document_type: str,
    user: UserContext = Depends(get_current_user),
) -> ApiResponse:
    authorize_cliente_access(cliente_id, user)
    registry = _read_registry(cliente_id)
    docs = registry.get("documents") if isinstance(registry.get("documents"), dict) else {}
    bucket = docs.get(document_type) if isinstance(docs, dict) else None
    versions = bucket.get("versions") if isinstance(bucket, dict) and isinstance(bucket.get("versions"), list) else []
    current = next((v for v in versions if isinstance(v, dict) and bool(v.get("is_current"))), None)
    if not isinstance(current, dict):
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="No hay version vigente para validar.")
    qc = _quality_check_version(current, cliente_id=cliente_id)
    return ApiResponse(
        data={
            "cliente_id": cliente_id,
            "document_type": document_type,
            "document_version": current.get("document_version"),
            "state": current.get("state"),
            "quality_check": qc,
        }
    )


@router.get("/{cliente_id}/documentos/{document_type}/evidence-gate", response_model=ApiResponse)
def get_document_evidence_gate(
    cliente_id: str,
    document_type: str,
    user: UserContext = Depends(get_current_user),
) -> ApiResponse:
    authorize_cliente_access(cliente_id, user)
    gate = _evidence_gate_for_current_version(cliente_id, document_type)
    return ApiResponse(data=gate)


@router.post("/{cliente_id}/documentos/{document_type}/estado", response_model=ApiResponse)
def post_document_state_transition(
    cliente_id: str,
    document_type: str,
    payload: dict[str, Any],
    user: UserContext = Depends(get_current_user),
) -> ApiResponse:
    authorize_cliente_access(cliente_id, user)
    target_state = _safe_text(payload.get("target_state")).lower()
    reason = _safe_text(payload.get("reason"))
    if not target_state:
        raise HTTPException(status_code=status.HTTP_422_UNPROCESSABLE_CONTENT, detail="target_state es obligatorio.")
    updated = _transition_document_state(
        cliente_id=cliente_id,
        document_type=document_type,
        target_state=target_state,
        changed_by=user.sub,
        changed_role=user.role,
        reason=reason,
    )
    _append_history(
        cliente_id,
        {
            "kind": "document_state_transition",
            "generated_at": datetime.now(timezone.utc).isoformat(),
            "document_type": document_type,
            "document_version": updated.get("document_version"),
            "state": updated.get("state"),
            "changed_by": user.sub,
            "changed_at": datetime.now(timezone.utc).isoformat(),
            "from_state": (updated.get("state_history") or [{}])[-1].get("from_state"),
            "to_state": (updated.get("state_history") or [{}])[-1].get("to_state"),
            "reason": reason,
            "status": "success",
        },
    )
    return ApiResponse(
        data={
            "cliente_id": cliente_id,
            "document_type": document_type,
            "updated_version": updated,
        }
    )


@router.post("/{cliente_id}/documentos/{document_type}/emitir", response_model=ApiResponse)
def post_document_issue(
    cliente_id: str,
    document_type: str,
    payload: dict[str, Any],
    user: UserContext = Depends(get_current_user),
) -> ApiResponse:
    authorize_cliente_access(cliente_id, user)
    reason = _safe_text(payload.get("reason")) or "Emision final"

    # 1) Transicion formal a issued (valida rol + flujo)
    updated = _transition_document_state(
        cliente_id=cliente_id,
        document_type=document_type,
        target_state="issued",
        changed_by=user.sub,
        changed_role=user.role,
        reason=reason,
    )

    artifacts = updated.get("artifacts") if isinstance(updated.get("artifacts"), list) else []
    docx_artifact = next(
        (
            a
            for a in artifacts
            if isinstance(a, dict) and _safe_text(a.get("artifact_type")) == "docx" and _safe_text(a.get("artifact_path"))
        ),
        None,
    )
    if not isinstance(docx_artifact, dict):
        raise HTTPException(
            status_code=status.HTTP_422_UNPROCESSABLE_CONTENT,
            detail="No existe artefacto DOCX para generar PDF de emisión.",
        )
    docx_path = Path(str(docx_artifact.get("artifact_path"))).resolve()
    if not docx_path.exists():
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail="No se encontró el archivo DOCX para emisión.",
        )

    # 2) Render PDF final con sello EMITIDO
    pdf_bytes = _emit_pdf_from_docx(docx_path=docx_path)
    timestamp = datetime.now(timezone.utc).strftime("%Y%m%dT%H%M%SZ")
    out_dir = REPORTS_META_DIR / cliente_id / "reportes"
    out_dir.mkdir(parents=True, exist_ok=True)
    pdf_path = out_dir / f"{document_type}_v{int(updated.get('document_version') or 0)}_issued_{timestamp}.pdf"
    pdf_path.write_bytes(pdf_bytes)
    pdf_hash = hashlib.sha256(pdf_bytes).hexdigest()

    # 3) Persistir en registro documental
    registry = _read_registry(cliente_id)
    docs = registry.get("documents") if isinstance(registry.get("documents"), dict) else {}
    bucket = docs.get(document_type) if isinstance(docs, dict) else None
    versions = bucket.get("versions") if isinstance(bucket, dict) and isinstance(bucket.get("versions"), list) else []
    current = next((v for v in versions if isinstance(v, dict) and bool(v.get("is_current"))), None)
    if isinstance(current, dict):
        current_artifacts = current.get("artifacts")
        if not isinstance(current_artifacts, list):
            current_artifacts = []
            current["artifacts"] = current_artifacts
        current_artifacts.append(
            {
                "artifact_type": "pdf",
                "artifact_path": str(pdf_path),
                "artifact_hash": pdf_hash,
                "template_version": _safe_text(
                    ((current.get("generation_metadata") or {}) if isinstance(current.get("generation_metadata"), dict) else {}).get(
                        "template_version"
                    )
                )
                or "v1",
                "size_bytes": len(pdf_bytes),
            }
        )
        current["issued_by"] = user.sub
        current["issued_at"] = datetime.now(timezone.utc).isoformat()
        current["seal"] = "EMITIDO"
        current["updated_at"] = datetime.now(timezone.utc).isoformat()
    _write_registry(cliente_id, registry)

    _append_history(
        cliente_id,
        {
            "kind": "document_issued",
            "generated_at": datetime.now(timezone.utc).isoformat(),
            "document_type": document_type,
            "document_version": updated.get("document_version"),
            "state": "issued",
            "changed_by": user.sub,
            "reason": reason,
            "pdf_path": str(pdf_path),
            "pdf_hash": pdf_hash,
            "size_bytes": len(pdf_bytes),
            "status": "success",
        },
    )

    return ApiResponse(
        data={
            "cliente_id": cliente_id,
            "document_type": document_type,
            "document_version": updated.get("document_version"),
            "state": "issued",
            "issued_by": user.sub,
            "issued_at": datetime.now(timezone.utc).isoformat(),
            "pdf_artifact_path": str(pdf_path),
            "pdf_artifact_hash": pdf_hash,
        }
    )
